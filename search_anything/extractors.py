"""Turn files into searchable text.

Each supported format has a small extractor. Anything else that looks like
plain text (no NUL bytes, valid UTF-8) is read as text too, so unusual code /
config files work without being listed. Binary files are still searchable by
name and folder.
"""

from __future__ import annotations

import html
import json
import logging
import os
import re
import threading
import zipfile
from email import policy
from email.parser import BytesParser
from pathlib import Path

log = logging.getLogger(__name__)

# Bump when extraction improves, so already-indexed files are read again once.
# v3: PDFs read with pdfium (much faster) and Thai text repaired.
EXTRACTOR_VERSION = 3

MAX_CONTENT_BYTES = 20 * 1024 * 1024  # don't read content of files bigger than this
MAX_TEXT_CHARS = 500_000  # cap stored text per file
MAX_PDF_PAGES = 300  # very long PDFs: first pages only (bounded memory)
SNIFF_MAX_BYTES = 2 * 1024 * 1024  # unknown extensions: sniff files up to this size

TEXT_EXTENSIONS = {
    # documents / data
    ".txt", ".text", ".md", ".markdown", ".mdx", ".rst", ".adoc", ".org", ".tex", ".bib", ".log",
    ".csv", ".tsv", ".json", ".jsonl", ".ndjson", ".geojson", ".yaml", ".yml", ".toml", ".ini",
    ".cfg", ".conf", ".properties", ".env", ".xml", ".svg", ".srt", ".vtt", ".sub", ".ics", ".vcf",
    ".css", ".scss", ".sass", ".less",
    # code
    ".py", ".pyi", ".js", ".mjs", ".cjs", ".ts", ".tsx", ".jsx", ".vue", ".svelte", ".astro",
    ".java", ".kt", ".kts", ".scala", ".groovy", ".gradle", ".go", ".rs", ".c", ".h", ".cc", ".cpp",
    ".hpp", ".cs", ".fs", ".vb", ".swift", ".m", ".mm", ".dart", ".php", ".rb", ".pl", ".pm", ".lua",
    ".r", ".jl", ".ex", ".exs", ".erl", ".hs", ".elm", ".clj", ".cljs", ".lisp", ".el", ".nim",
    ".zig", ".sol", ".sql", ".graphql", ".gql", ".proto", ".tf", ".hcl", ".nix",
    ".sh", ".bash", ".zsh", ".fish", ".ps1", ".psm1", ".bat", ".cmd", ".vbs", ".ahk",
}
TEXT_FILENAMES = {
    "dockerfile", "makefile", "gnumakefile", "cmakelists.txt", "procfile", "gemfile", "rakefile",
    "vagrantfile", "jenkinsfile", "readme", "license", "licence", "changelog", "authors", "notice",
}
HTML_EXTENSIONS = {".html", ".htm", ".xhtml"}
ODF_EXTENSIONS = {".odt", ".ods", ".odp", ".odg"}

# Shown in the UI / README.
SUPPORTED = {
    "ข้อความและโค้ด": "txt, md, csv, json, yaml, xml, log, โค้ดกว่า 80 ภาษา, Dockerfile, Makefile และไฟล์ข้อความอื่นๆ",
    "เอกสาร": "pdf, docx, rtf, odt, epub, html",
    "ตาราง": "xlsx, xlsm, ods, csv, tsv",
    "สไลด์": "pptx, odp",
    "อื่นๆ": "eml (อีเมล), ipynb (Jupyter), zip (รายชื่อไฟล์ข้างใน)",
    "ภาพสแกน (OCR)": "PDF ที่สแกนมา และรูปภาพ jpg, png, tiff (ตั้งค่าในหัวข้อ OCR) — ต้องมี Tesseract",
}


# pdfium may only be used by one thread at a time (in the whole process).
PDFIUM_LOCK = threading.Lock()

# ------------------------------------------------------------------ helpers
# Thai fonts in PDFs often use private-use code points for shifted tone marks
# and vowels; map them back to the real characters.
_THAI_PUA = str.maketrans({
    0xF700: "\u0e10", 0xF701: "\u0e34", 0xF702: "\u0e35", 0xF703: "\u0e36", 0xF704: "\u0e37",
    0xF705: "\u0e48", 0xF706: "\u0e49", 0xF707: "\u0e4a", 0xF708: "\u0e4b", 0xF709: "\u0e4c",
    0xF70A: "\u0e48", 0xF70B: "\u0e49", 0xF70C: "\u0e4a", 0xF70D: "\u0e4b", 0xF70E: "\u0e4c",
    0xF70F: "\u0e0d", 0xF710: "\u0e31", 0xF711: "\u0e4d", 0xF712: "\u0e47", 0xF713: "\u0e48",
    0xF714: "\u0e49", 0xF715: "\u0e4a", 0xF716: "\u0e4b", 0xF717: "\u0e4c", 0xF718: "\u0e38",
    0xF719: "\u0e39", 0xF71A: "\u0e3a",
})
_SARA_AM_TONE = re.compile("\u0e4d([\u0e48-\u0e4b])\u0e32")  # ํ + tone + า -> tone + ำ
_SARA_AM = re.compile("\u0e4d\u0e32")  # ํ + า -> ำ
_GAP_AFTER_AM = re.compile("\u0e33 (?=[\u0e01-\u0e2e\u0e40-\u0e44])")  # "สำ นักงาน" (PDF layout gap)
_THAI = re.compile("[\u0e00-\u0e7f\uf700-\uf71a]")


def normalize_thai(text: str, pdf: bool = False) -> str:
    """Make Thai text match what people type: real tone marks, one-character ำ."""
    if not text or not _THAI.search(text):
        return text
    text = text.translate(_THAI_PUA)
    text = _SARA_AM.sub("\u0e33", _SARA_AM_TONE.sub("\\1\u0e33", text))
    if pdf:
        text = _GAP_AFTER_AM.sub("\u0e33", text)
    return text


def decode(raw: bytes) -> str:
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16", errors="replace")
    for encoding in ("utf-8-sig", "cp874", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


_BLOCK_TAGS = re.compile(r"<\s*(br|/p|/div|/li|/tr|/h[1-6]|/section|/article|/title|/blockquote)\b[^>]*>", re.I)
_DROP = re.compile(r"<(script|style|head)\b.*?</\1\s*>", re.I | re.S)
_TAG = re.compile(r"<[^>]+>")


def html_to_text(markup: str) -> str:
    text = _DROP.sub(" ", markup)
    text = _BLOCK_TAGS.sub("\n", text)
    text = html.unescape(_TAG.sub("", text))
    return re.sub(r"[ \t]+\n", "\n", re.sub(r"\n{3,}", "\n\n", text)).strip()


class _Collector:
    """Gathers text pieces and stops once the size cap is reached."""

    def __init__(self):
        self.parts: list[str] = []
        self.size = 0

    @property
    def full(self) -> bool:
        return self.size >= MAX_TEXT_CHARS

    def add(self, text: str) -> None:
        if text and not self.full:
            self.parts.append(text)
            self.size += len(text) + 1

    def text(self) -> str:
        return "\n".join(self.parts)[:MAX_TEXT_CHARS]


# --------------------------------------------------------------- extractors
def _plain(path: Path) -> str:
    raw = path.read_bytes()
    if b"\x00" in raw[:4096] and not raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return ""
    return decode(raw)


def _html(path: Path) -> str:
    return html_to_text(decode(path.read_bytes()))


def _pdf(path: Path) -> str:
    try:
        text = _pdf_pdfium(path)
    except ImportError:
        text = _pdf_pypdf(path)
    return normalize_thai(text, pdf=True)


def _pdf_pdfium(path: Path) -> str:
    import pypdfium2 as pdfium

    out = _Collector()
    with PDFIUM_LOCK:
        pdf = pdfium.PdfDocument(str(path))
        try:
            for i in range(min(len(pdf), MAX_PDF_PAGES)):
                if out.full:
                    break
                page = pdf[i]
                try:
                    textpage = page.get_textpage()
                    out.add(textpage.get_text_range().replace("\r\n", "\n"))
                    textpage.close()
                finally:
                    page.close()
        finally:
            pdf.close()
    return out.text()


def _pdf_pypdf(path: Path) -> str:
    from pypdf import PdfReader

    out = _Collector()
    for i, page in enumerate(PdfReader(str(path)).pages):
        if i >= MAX_PDF_PAGES or out.full:
            break
        out.add(page.extract_text() or "")
    return out.text()


def _docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    out = _Collector()
    for p in document.paragraphs:
        out.add(p.text)
    for table in document.tables:
        for row in table.rows:
            out.add("\t".join(cell.text for cell in row.cells))
    return out.text()


def _xlsx(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    out = _Collector()
    try:
        for ws in wb.worksheets:
            out.add(f"# {ws.title}")
            for row in ws.iter_rows(values_only=True):
                cells = [str(v) for v in row if v is not None and str(v) != ""]
                if cells:
                    out.add("\t".join(cells))
                if out.full:
                    break
    finally:
        wb.close()
    return out.text()


def _pptx(path: Path) -> str:
    from pptx import Presentation

    out = _Collector()
    for n, slide in enumerate(Presentation(str(path)).slides, start=1):
        out.add(f"# สไลด์ {n}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.add(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    out.add("\t".join(cell.text for cell in row.cells))
        if slide.has_notes_slide:
            out.add(slide.notes_slide.notes_text_frame.text)
    return out.text()


def _odf(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        xml = z.read("content.xml").decode("utf-8", errors="replace")
    xml = re.sub(r"<text:(line-break|tab)\s*/>", lambda m: "\n" if m.group(1) == "line-break" else "\t", xml)
    xml = re.sub(r"</(text:p|text:h|table:table-row)>", "\n", xml)
    xml = re.sub(r"</table:table-cell>", "\t", xml)
    return html.unescape(_TAG.sub("", xml)).strip()


def _epub(path: Path) -> str:
    out = _Collector()
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        order = []
        # Follow the reading order (spine) from the package file when possible.
        opf = next((n for n in names if n.endswith(".opf")), None)
        if opf:
            pkg = z.read(opf).decode("utf-8", errors="replace")
            base = opf.rsplit("/", 1)[0] + "/" if "/" in opf else ""
            items = dict(re.findall(r'<item\b[^>]*\bid="([^"]+)"[^>]*\bhref="([^"]+)"', pkg))
            items.update({i: h for h, i in re.findall(r'<item\b[^>]*\bhref="([^"]+)"[^>]*\bid="([^"]+)"', pkg)})
            order = [base + items[i] for i in re.findall(r'<itemref\b[^>]*\bidref="([^"]+)"', pkg) if i in items]
        order = [n for n in order if n in names] or sorted(n for n in names if n.lower().endswith((".xhtml", ".html", ".htm")))
        for name in order:
            out.add(html_to_text(z.read(name).decode("utf-8", errors="replace")))
            if out.full:
                break
    return out.text()


def _rtf(path: Path) -> str:
    raw = path.read_bytes().decode("latin-1")
    m = re.search(r"\\ansicpg(\d+)", raw)
    codepage = f"cp{m.group(1)}" if m else "cp1252"
    # Drop destinations we never want to show ({\*\...}, fonts, colours, pictures …).
    raw = re.sub(r"\{\\\*[^{}]*(\{[^{}]*\}[^{}]*)*\}", "", raw)
    for group in ("fonttbl", "colortbl", "stylesheet", "info", "pict"):
        raw = re.sub(r"\{\\" + group + r"[^{}]*(\{[^{}]*\}[^{}]*)*\}", "", raw)
    out, pending = [], bytearray()

    def flush():
        if pending:
            out.append(pending.decode(codepage, errors="replace"))
            pending.clear()

    i = 0
    token = re.compile(r"\\'([0-9a-fA-F]{2})|\\u(-?\d+)\??|\\(par|line|tab)\b ?|\\[a-zA-Z]+-?\d* ?|\\([\\{}])|[{}]|\r|\n|([^\\{}\r\n]+)")
    for t in token.finditer(raw):
        hexbyte, uni, brk, esc, text = t.group(1), t.group(2), t.group(3), t.group(4), t.group(5)
        if hexbyte:
            pending.append(int(hexbyte, 16))
            continue
        flush()
        if uni:
            out.append(chr(int(uni) % 65536))
        elif brk:
            out.append("\t" if brk == "tab" else "\n")
        elif esc:
            out.append(esc)
        elif text:
            out.append(text)
        i += 1
    flush()
    return re.sub(r"\n{3,}", "\n\n", "".join(out)).strip()


def _eml(path: Path) -> str:
    msg = BytesParser(policy=policy.default).parse(path.open("rb"))
    out = _Collector()
    for header in ("Subject", "From", "To", "Cc", "Date"):
        if msg[header]:
            out.add(f"{header}: {msg[header]}")
    body = msg.get_body(preferencelist=("plain", "html"))
    if body is not None:
        content = body.get_content()
        out.add(html_to_text(content) if body.get_content_type() == "text/html" else content)
    for part in msg.iter_attachments():
        if part.get_filename():
            out.add(f"ไฟล์แนบ: {part.get_filename()}")
    return out.text()


def _ipynb(path: Path) -> str:
    nb = json.loads(path.read_bytes())
    out = _Collector()
    for cell in nb.get("cells", []):
        src = cell.get("source", "")
        out.add("".join(src) if isinstance(src, list) else src)
        for o in cell.get("outputs", []):
            text = o.get("text") or (o.get("data") or {}).get("text/plain") or ""
            out.add("".join(text) if isinstance(text, list) else text)
    return out.text()


def _zip(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        return "\n".join(n for n in z.namelist() if not n.endswith("/"))[:MAX_TEXT_CHARS]


# ---------------------------------------------------------- native (Rust)
# The optional `sa_native` module (native/ in this repo, written in Rust) gives
# exactly the same text for Word/Excel/PowerPoint files, 6-30x faster and
# without holding the GIL, so scan threads really run in parallel. Without it
# (or with SA_NATIVE=0) the Python readers above are used. If the Rust reader
# fails on a file, the Python reader gets a try too.
try:
    if os.environ.get("SA_NATIVE", "1") == "0":
        raise ImportError("disabled by SA_NATIVE=0")
    import sa_native as _native
except ImportError:
    _native = None

NATIVE = _native is not None


def _with_native(native_fn, python_fn):
    def handler(path: Path) -> str:
        try:
            return native_fn(str(path), MAX_TEXT_CHARS)
        except Exception as exc:
            log.debug("native reader failed for %s (%s); using Python", path, exc)
            return python_fn(path)

    handler.__name__ = python_fn.__name__ + "_native"
    return handler


if _native is not None:
    _docx_fast = _with_native(_native.docx_text, _docx)
    _xlsx_fast = _with_native(_native.xlsx_text, _xlsx)
    _pptx_fast = _with_native(_native.pptx_text, _pptx)
else:
    _docx_fast, _xlsx_fast, _pptx_fast = _docx, _xlsx, _pptx

HANDLERS = {
    ".pdf": _pdf, ".docx": _docx_fast, ".xlsx": _xlsx_fast, ".xlsm": _xlsx_fast, ".pptx": _pptx_fast,
    ".rtf": _rtf, ".epub": _epub, ".eml": _eml, ".ipynb": _ipynb, ".zip": _zip,
    **{e: _html for e in HTML_EXTENSIONS}, **{e: _odf for e in ODF_EXTENSIONS},
}


OFFICE_EXTENSIONS = {".docx", ".xlsx", ".xlsm", ".pptx"}


def is_heavy(path: str) -> bool:
    """Readers worth a helper process: pdfium (one at a time per process), and
    the Python Office readers when the Rust module is missing (they hold the GIL)."""
    suffix = os.path.splitext(path)[1].lower()
    return suffix == ".pdf" or (not NATIVE and suffix in OFFICE_EXTENSIONS)


def is_plain_text_type(path: Path) -> bool:
    """Files whose text on disk is exactly what we index (line numbers match)."""
    return path.suffix.lower() in TEXT_EXTENSIONS or path.name.lower() in TEXT_FILENAMES


def extract_text(path: Path, size: int | None = None) -> str:
    """Return the searchable text content of a file ('' if unsupported)."""
    suffix = path.suffix.lower()
    try:
        if size is None:
            size = path.stat().st_size
        if size > MAX_CONTENT_BYTES:
            return ""
        handler = HANDLERS.get(suffix)
        if handler:
            return handler(path)[:MAX_TEXT_CHARS]
        if is_plain_text_type(path):
            return _plain(path)[:MAX_TEXT_CHARS]
        if size <= SNIFF_MAX_BYTES:
            # Unknown extension: accept it if it is clearly text.
            raw = path.read_bytes()
            if raw and b"\x00" not in raw[:8192]:
                try:
                    return raw.decode("utf-8")[:MAX_TEXT_CHARS]
                except UnicodeDecodeError:
                    pass
    except Exception as exc:  # corrupt / locked / unreadable / password-protected files
        log.debug("could not extract %s: %s", path, exc)
    return ""

import time
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from search_anything.app import Config, create_app
from search_anything.db import Database
from search_anything.indexer import Indexer

FAST_SCRYPT = 2**10  # keep tests fast; production uses 2**15


def make_client(tmp_path: Path, **kw) -> TestClient:
    config = Config(
        data_dir=tmp_path / "data",
        allowed_hosts=["testserver"],
        scrypt_n=FAST_SCRYPT,
        background_indexing=False,
        **kw,
    )
    return TestClient(create_app(config))


@pytest.fixture
def docs(tmp_path: Path) -> Path:
    root = tmp_path / "docs"
    (root / "sub").mkdir(parents=True)
    (root / "สัญญาเช่าบ้าน.txt").write_text("ค่าเช่าเดือนละ 8000 บาท จ่ายทุกวันที่ 5", encoding="utf-8")
    (root / "sub" / "invoice-2026.md").write_text("Invoice for ACME Corp\nTotal: 1200 USD")
    (root / "photo.jpg").write_bytes(b"\xff\xd8\xff\x00binary")
    (root / "node_modules").mkdir()
    (root / "node_modules" / "skip.txt").write_text("ACME should not be indexed here")
    return root


def ids(resp) -> list[str]:
    return [r["title"] for r in resp.json()["results"]]


# ------------------------------------------------------------------ indexer
def test_indexer_names_content_thai_and_incremental(tmp_path, docs):
    db = Database(tmp_path / "t.db")
    indexer = Indexer(db, [str(docs)])
    stats = indexer.run()
    assert stats.added_or_updated == 3  # node_modules skipped

    assert [r["title"] for r in db.search("ACME")] == ["invoice-2026.md"]
    assert [r["title"] for r in db.search("ค่าเช่า")] == ["สัญญาเช่าบ้าน.txt"]  # Thai content
    assert [r["title"] for r in db.search("สัญญา")] == ["สัญญาเช่าบ้าน.txt"]  # Thai file name
    assert [r["title"] for r in db.search("photo")] == ["photo.jpg"]  # name-only file
    assert [r["title"] for r in db.search("sub")] == ["invoice-2026.md"]  # folder in path

    # Nothing changed -> nothing re-read.
    assert indexer.run().added_or_updated == 0

    # Edit, delete.
    (docs / "sub" / "invoice-2026.md").write_text("Paid in full by Globex")
    (docs / "photo.jpg").unlink()
    stats = indexer.run()
    assert stats.added_or_updated == 1 and stats.removed == 1
    assert db.search("ACME") == []
    assert [r["title"] for r in db.search("Globex")] == ["invoice-2026.md"]


def test_search_short_terms_and_special_chars(tmp_path):
    db = Database(":memory:")
    db.add_item("note", "Wi-Fi บ้าน", "ssid: home_5g")
    db.add_item("note", "C++ notes", 'quote " and * stars')
    assert [r["title"] for r in db.search("wi")] == ["Wi-Fi บ้าน"]  # < 3 chars uses LIKE
    assert [r["title"] for r in db.search("C++")] == ["C++ notes"]
    assert [r["title"] for r in db.search('" and *')] == ["C++ notes"]
    assert [r["title"] for r in db.search("home_5g บ้")] == ["Wi-Fi บ้าน"]  # mixed long + short
    assert db.search("%") == []


# ---------------------------------------------------------------- local mode
def test_notes_bookmarks_crud_and_search(tmp_path):
    c = make_client(tmp_path)
    note = c.post("/api/items", json={"kind": "note", "title": "เลขบัญชี", "content": "กสิกร 123-4-56789", "tags": ["bank"]}).json()
    c.post("/api/items", json={"kind": "bookmark", "title": "Docs", "url": "https://fastapi.tiangolo.com", "content": "python web"})

    assert ids(c.get("/api/search", params={"q": "กสิกร"})) == ["เลขบัญชี"]
    assert ids(c.get("/api/search", params={"q": "#bank"})) == ["เลขบัญชี"]
    assert ids(c.get("/api/search", params={"q": "tiangolo"})) == ["Docs"]
    assert ids(c.get("/api/search", params={"q": "python", "kind": "note"})) == []

    c.put(f"/api/items/{note['id']}", json={"kind": "note", "title": "เลขบัญชี", "content": "SCB 999"})
    assert ids(c.get("/api/search", params={"q": "กสิกร"})) == []
    assert c.get(f"/api/items/{note['id']}").json()["content"] == "SCB 999"

    assert c.delete(f"/api/items/{note['id']}").json() == {"ok": True}
    assert c.get(f"/api/items/{note['id']}").status_code == 404


def test_password_vault(tmp_path):
    c = make_client(tmp_path)
    item = {"kind": "password", "title": "Gmail", "username": "me@example.com", "url": "mail.google.com", "password": "s3cret-pass!"}
    assert c.post("/api/items", json=item).status_code == 423  # no vault yet

    assert c.post("/api/vault/setup", json={"master_password": "short"}).status_code == 400
    assert c.post("/api/vault/setup", json={"master_password": "correct horse"}).status_code == 200
    created = c.post("/api/items", json=item).json()
    pid = created["id"]

    # Searchable by metadata, never by the secret, and the secret is not in any response.
    assert ids(c.get("/api/search", params={"q": "gmail"})) == ["Gmail"]
    assert ids(c.get("/api/search", params={"q": "example.com"})) == ["Gmail"]
    assert ids(c.get("/api/search", params={"q": "s3cret"})) == []
    assert "s3cret" not in c.get(f"/api/items/{pid}").text
    assert "s3cret" not in c.get("/api/search", params={"q": "gmail"}).text

    # Not stored in plaintext on disk.
    raw = (tmp_path / "data" / "search.db").read_bytes() + (tmp_path / "data" / "search.db-wal").read_bytes()
    assert b"s3cret-pass!" not in raw

    assert c.post(f"/api/items/{pid}/reveal").json() == {"password": "s3cret-pass!"}
    c.post("/api/vault/lock")
    assert c.post(f"/api/items/{pid}/reveal").status_code == 423
    assert c.post("/api/vault/unlock", json={"master_password": "wrong one"}).status_code == 401
    assert c.post("/api/vault/unlock", json={"master_password": "correct horse"}).status_code == 200

    # Update without a new password keeps the old one.
    c.put(f"/api/items/{pid}", json={**item, "password": None, "title": "Gmail personal"})
    assert c.post(f"/api/items/{pid}/reveal").json() == {"password": "s3cret-pass!"}


def test_local_mode_folders_and_host_check(tmp_path, docs):
    c = make_client(tmp_path)
    assert c.put("/api/folders", json={"folders": [str(tmp_path / "nope")]}).status_code == 400
    assert c.put("/api/folders", json={"folders": [str(docs)]}).status_code == 200
    c.app.state.indexer.run()
    assert ids(c.get("/api/search", params={"q": "ACME"})) == ["invoice-2026.md"]

    # Folders are remembered across restarts.
    c2 = make_client(tmp_path)
    assert c2.get("/api/status").json()["folders"] == [str(docs.resolve())]

    # Requests with a foreign Host header (DNS rebinding) are rejected.
    assert c.get("/api/status", headers={"host": "evil.example"}).status_code == 400


# --------------------------------------------------------------- server mode
def test_server_mode_requires_login_and_supports_upload(tmp_path):
    with pytest.raises(ValueError):
        make_client(tmp_path, mode="server")

    c = make_client(tmp_path, mode="server", password="hunter22")
    assert c.get("/api/status").json() == {"mode": "server", "authenticated": False}
    assert c.get("/api/search", params={"q": "x"}).status_code == 401
    assert c.post("/api/items", json={"kind": "note", "title": "x"}).status_code == 401
    assert c.put("/api/folders", json={"folders": []}).status_code in (401, 403)

    assert c.post("/api/login", json={"password": "nope"}).status_code == 401
    assert c.post("/api/login", json={"password": "hunter22"}).status_code == 200
    assert c.get("/api/status").json()["authenticated"] is True

    r = c.post("/api/upload", files=[("files", ("../../ไฟล์ทดสอบ.txt", "รายงานประจำปี".encode())), ("files", ("ไฟล์ทดสอบ.txt", b"dup"))])
    assert r.json()["saved"] == ["ไฟล์ทดสอบ.txt", "ไฟล์ทดสอบ (1).txt"]  # no path traversal, no overwrite
    c.app.state.indexer.run()
    results = c.get("/api/search", params={"q": "รายงาน"}).json()["results"]
    assert [r["title"] for r in results] == ["ไฟล์ทดสอบ.txt"]
    dl = c.get(f"/api/files/{results[0]['id']}/download")
    assert dl.content == "รายงานประจำปี".encode()
    assert c.post(f"/api/files/{results[0]['id']}/open").status_code == 403

    c.post("/api/logout")
    assert c.get("/api/search").status_code == 401


def test_ui_served(tmp_path):
    c = make_client(tmp_path)
    assert "Search Anything" in c.get("/").text
    assert c.get("/static/app.js").status_code == 200


def test_trigger_during_scan_is_not_dropped(tmp_path, docs):
    db = Database(tmp_path / "t.db")
    indexer = Indexer(db, [str(docs)])
    with indexer._run_lock:  # simulate a scan in progress
        indexer.trigger()
        indexer.trigger()  # coalesced with the queued one
        (docs / "late.txt").write_text("zebracorn")
    # The queued background scan must run after the current one finishes.
    deadline = time.time() + 5
    while not db.search("zebracorn") and time.time() < deadline:
        time.sleep(0.05)
    assert [r["title"] for r in db.search("zebracorn")] == ["late.txt"]


# ------------------------------------------------------------------- upload
def test_upload_local_mode_search_and_delete(tmp_path, docs):
    c = make_client(tmp_path, folders=[str(docs)])
    c.app.state.indexer.run()
    r = c.post("/api/upload", files=[("files", ("ใบเสนอราคา.txt", "ราคารวม 45,000 บาท".encode()))])
    assert r.status_code == 200
    item = r.json()["items"][0]
    assert item["uploaded"] is True
    # Searchable as soon as the request returns, no rescan needed.
    assert ids(c.get("/api/search", params={"q": "ราคารวม"})) == ["ใบเสนอราคา.txt"]
    # A rescan keeps it (the uploads folder is always indexed) …
    c.app.state.indexer.run()
    assert ids(c.get("/api/search", params={"q": "ราคารวม"})) == ["ใบเสนอราคา.txt"]
    # … and it is not listed as a user folder.
    assert c.get("/api/status").json()["folders"] == [str(docs.resolve())]

    # Uploaded files can be deleted; files from the user's own folders cannot.
    other = c.get("/api/search", params={"q": "ACME"}).json()["results"][0]
    assert other["uploaded"] is False
    assert c.delete(f"/api/items/{other['id']}").status_code == 400
    assert (docs / "sub" / "invoice-2026.md").exists()
    assert c.delete(f"/api/items/{item['id']}").json() == {"ok": True}
    assert ids(c.get("/api/search", params={"q": "ราคารวม"})) == []
    assert not (tmp_path / "data" / "uploads" / "ใบเสนอราคา.txt").exists()


def test_upload_size_limit_and_windows_style_names(tmp_path):
    c = make_client(tmp_path, max_upload_mb=1)
    big = b"x" * (1024 * 1024 + 1)
    assert c.post("/api/upload", files=[("files", ("big.txt", big))]).status_code == 413
    assert not list((tmp_path / "data" / "uploads").iterdir())  # partial file removed
    r = c.post("/api/upload", files=[("files", ("C:\\Users\\me\\..\\.secret.txt", b"hi"))])
    assert r.json()["saved"] == ["secret.txt"]


# ------------------------------------------------------------------ ranking
def test_name_matches_rank_before_content_matches(tmp_path):
    root = tmp_path / "r"
    root.mkdir()
    (root / "budget-2026.txt").write_text("numbers")
    (root / "notes.txt").write_text("the budget is tight")
    db = Database(tmp_path / "t.db")
    Indexer(db, [str(root)]).run()
    db.add_item("note", "Budget meeting", "agenda")
    db.add_item("note", "Groceries", "keep the budget low")
    titles = [r["title"] for r in db.search("budget")]
    assert set(titles[:2]) == {"budget-2026.txt", "Budget meeting"}
    assert set(titles[2:]) == {"notes.txt", "Groceries"}
    snippet_row = next(r for r in db.search("tight") if r["title"] == "notes.txt")
    assert "budget is tight" in snippet_row["body"]


def test_old_single_file_database_is_migrated(tmp_path):
    import sqlite3, time as _t
    path = tmp_path / "search.db"
    Database(path).close()
    conn = sqlite3.connect(path)
    conn.execute("INSERT INTO items(kind,title,body,path,created_at,updated_at) VALUES('file','old.txt','x','/old.txt',?,?)", (_t.time(), _t.time()))
    conn.execute("UPDATE meta SET value='1' WHERE key='schema'")
    conn.commit(); conn.close()
    db = Database(path)
    assert db.counts()["file"] == 0 and db.search("old.txt") == []


def test_sorting(tmp_path):
    root = tmp_path / "s"
    root.mkdir()
    import os as _os
    for i, (name, size) in enumerate([("beta.txt", 30), ("alpha.txt", 10), ("gamma.txt", 20)]):
        p = root / name
        p.write_text("kiwi " + "x" * size)
        _os.utime(p, (1_000_000 + i, 1_000_000 + i))  # beta oldest, gamma newest
    c = make_client(tmp_path, folders=[str(root)])
    c.app.state.indexer.run()
    c.post("/api/items", json={"kind": "note", "title": "Kiwi note", "content": "kiwi"})
    q = lambda sort, kind=None: ids(c.get("/api/search", params={"q": "kiwi", "sort": sort, **({"kind": kind} if kind else {})}))
    assert q("name", "file") == ["alpha.txt", "beta.txt", "gamma.txt"]
    assert q("name_desc", "file") == ["gamma.txt", "beta.txt", "alpha.txt"]
    assert q("newest", "file") == ["gamma.txt", "alpha.txt", "beta.txt"]
    assert q("oldest", "file") == ["beta.txt", "alpha.txt", "gamma.txt"]
    assert q("largest", "file") == ["beta.txt", "gamma.txt", "alpha.txt"]
    assert q("type")[:3] == ["alpha.txt", "beta.txt", "gamma.txt"] and q("type")[-1] == "Kiwi note"
    assert q("newest")[0] == "Kiwi note"  # just created
    assert c.get("/api/search", params={"q": "kiwi", "sort": "random"}).status_code == 422


def test_lines_for_viewer(tmp_path):
    root = tmp_path / "v"
    root.mkdir()
    (root / "code.py").write_text("import os\n\ndef kiwi():\n    return 'Kiwi kiwi'\n")
    import docx as _docx
    d = _docx.Document(); d.add_paragraph("หน้าแรก"); d.add_paragraph("kiwi ในเอกสาร"); d.save(root / "doc.docx")
    c = make_client(tmp_path, folders=[str(root)])
    c.app.state.indexer.run()
    by_title = {r["title"]: r for r in c.get("/api/search", params={"q": "kiwi"}).json()["results"]}
    code = c.get(f"/api/items/{by_title['code.py']['id']}/lines").json()
    assert code["source"] == "disk" and code["total_lines"] == 4
    assert code["lines"][2] == "def kiwi():" and code["lines"][1] == ""  # exact lines, blank kept
    doc = c.get(f"/api/items/{by_title['doc.docx']['id']}/lines").json()
    assert doc["source"] == "index" and doc["lines"] == ["หน้าแรก", "kiwi ในเอกสาร"]
    note = c.post("/api/items", json={"kind": "note", "title": "n", "content": "a\nkiwi b"}).json()
    assert c.get(f"/api/items/{note['id']}/lines").json()["lines"] == ["a", "kiwi b"]
    bm = c.post("/api/items", json={"kind": "bookmark", "title": "b", "url": "x.com"}).json()
    assert c.get(f"/api/items/{bm['id']}/lines").status_code == 400
    assert "เอกสาร" in c.get("/api/status").json()["supported_formats"]


def test_all_view_shows_files_before_searching(tmp_path):
    import os as _os
    import time as _t
    root = tmp_path / "home"
    root.mkdir()
    (root / "old-report.txt").write_text("x")
    _os.utime(root / "old-report.txt", (1_000_000, 1_000_000))
    (root / "fresh-plan.md").write_text("y")  # modified now
    c = make_client(tmp_path, folders=[str(root)])
    c.app.state.indexer.run()
    _t.sleep(0.01)
    c.post("/api/items", json={"kind": "note", "title": "newest note"})
    titles = ids(c.get("/api/search", params={"q": ""}))
    assert titles == ["newest note", "fresh-plan.md", "old-report.txt"]  # files included, newest first
    assert ids(c.get("/api/search", params={"q": "", "kind": "note"})) == ["newest note"]
    assert ids(c.get("/api/search", params={"q": "", "kind": "file"})) == ["fresh-plan.md", "old-report.txt"]

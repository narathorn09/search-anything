"""The Rust readers (cahon_native) must give exactly the text the Python readers give.

Each test builds a document with the tricky parts of its format, then compares
the two readers character by character. Skipped when cahon_native is not built
(`pip install ./native`, needs Rust).
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import pytest

from cahon import extractors as ex

native = pytest.importorskip("cahon_native")
CAP = ex.MAX_TEXT_CHARS


def same(py_fn, native_fn, path: Path) -> str:
    expected = py_fn(path)[:CAP]
    got = native_fn(str(path), CAP)
    assert got == expected
    return got


def test_docx_matches_python_docx(tmp_path):
    import docx
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    d = docx.Document()
    d.add_heading("หัวข้อรายงาน สำนักงาน", 1)
    p = d.add_paragraph("ข้อความ ")
    p.add_run("ตัวหนา").bold = True
    p.add_run("\tหลังแท็บ")
    run = p.add_run("ก่อนขึ้นบรรทัด")
    run.add_break()  # line break -> "\n"
    run.add_text("หลังขึ้นบรรทัด")
    page = p.add_run("x")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")  # page break -> ""
    page._r.append(br)
    # A hyperlink with a run inside.
    link = OxmlElement("w:hyperlink")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "ลิงก์ไปเว็บ"
    r.append(t)
    link.append(r)
    p._p.append(link)
    d.add_paragraph("")  # empty paragraph
    # Merged cells: horizontal and vertical.
    table = d.add_table(rows=3, cols=3)
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            cell.text = f"r{i}c{j}"
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 2).merge(table.cell(2, 2))
    table.cell(2, 0).text = "สองย่อหน้า"
    table.cell(2, 0).add_paragraph("ย่อหน้าที่สอง")
    d.add_paragraph("หลังตาราง")
    path = tmp_path / "doc.docx"
    d.save(path)
    text = same(ex._docx, native.docx_text, path)
    assert "ลิงก์ไปเว็บ" in text and "r1c2" in text


def test_xlsx_matches_openpyxl(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ข้อมูล"
    ws.append(["ชื่อ", "จำนวน", "ราคา", "วันที่", "จริงไหม"])
    ws.append(["กาแฟ", 3, 45.5, dt.datetime(2024, 1, 31, 8, 30), True])
    ws.append(["ชา", 10, 0.1 + 0.2, dt.date(2023, 12, 25), False])
    ws.append([None, None, 1e-7, dt.time(14, 15), None])
    ws.append([None, 12345678901, 2.5e20, None, ""])
    ws["A7"] = "=1/0"  # formula without cached value -> empty
    ws["B8"] = 1234.5
    ws["C8"] = -0.0001
    ws["D8"] = dt.timedelta(days=1, hours=2)
    ws["D8"].number_format = "[h]:mm:ss"
    ws2 = wb.create_sheet("Sheet two")
    for i in range(50):
        ws2.append([f"แถว {i}", i * 1.25, None, "x" * (i % 3)])
    wb.create_sheet("ว่างเปล่า")
    path = tmp_path / "book.xlsx"
    wb.save(path)
    text = same(ex._xlsx, native.xlsx_text, path)
    assert "# ข้อมูล" in text and "2024-01-31 08:30:00" in text and "45.5" in text


def test_pptx_matches_python_pptx(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "สไลด์แรก งบประมาณ"
    body = s1.placeholders[1].text_frame
    body.text = "ข้อแรก"
    para = body.add_paragraph()
    para.text = "บรรทัด\vขึ้นใหม่ในย่อหน้า"  # soft line break -> "\v"
    s1.notes_slide.notes_text_frame.text = "โน้ตผู้บรรยาย"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "ตาราง"
    tbl = s2.shapes.add_table(2, 3, Inches(1), Inches(2), Inches(6), Inches(1)).table
    for i in range(2):
        for j in range(3):
            tbl.cell(i, j).text = f"ช่อง {i},{j}"
    box = s2.shapes.add_textbox(Inches(1), Inches(4), Inches(3), Inches(1))
    box.text_frame.text = "กล่องข้อความ"
    group = s2.shapes.add_group_shape()
    inner = group.shapes.add_textbox(Inches(1), Inches(5), Inches(2), Inches(1))
    inner.text_frame.text = "ในกลุ่ม (python-pptx ไม่อ่าน)"
    prs.slides.add_slide(prs.slide_layouts[6])  # blank slide
    path = tmp_path / "deck.pptx"
    prs.save(path)
    text = same(ex._pptx, native.pptx_text, path)
    assert "# สไลด์ 3" in text and "โน้ตผู้บรรยาย" in text and "ช่อง 1,2" in text


def test_text_cap_is_the_same(tmp_path, monkeypatch):
    import docx

    d = docx.Document()
    for i in range(300):
        d.add_paragraph(f"ย่อหน้าที่ {i} " + "ข" * 50)
    path = tmp_path / "long.docx"
    d.save(path)
    monkeypatch.setattr(ex, "MAX_TEXT_CHARS", 1000)
    expected = ex._docx(path)[:1000]
    assert native.docx_text(str(path), 1000) == expected
    assert len(expected) == 1000


def test_broken_files_fall_back_to_python(tmp_path):
    bad = tmp_path / "broken.docx"
    bad.write_bytes(b"not a zip at all")
    with pytest.raises(ValueError):
        native.docx_text(str(bad), CAP)
    assert ex.extract_text(bad) == ""  # both readers fail: no content, no crash

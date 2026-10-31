import json
import zipfile
from pathlib import Path

import pytest

from search_anything.db import Database
from search_anything.extractors import extract_text
from search_anything.indexer import Indexer


def make_xlsx(p: Path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "งบประมาณ"
    ws.append(["รายการ", "ยอด"])
    ws.append(["ค่าเช่าสำนักงาน", 45000])
    wb.save(p)


def make_pptx(p: Path):
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "แผนการตลาด Q3"
    slide.placeholders[1].text = "เปิดตัวสินค้า quokka"
    slide.notes_slide.notes_text_frame.text = "speaker note ลับ"
    prs.save(p)


def make_docx_with_table(p: Path):
    import docx
    d = docx.Document()
    d.add_paragraph("สัญญาจ้าง")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "ค่าจ้างรายเดือน"
    t.rows[0].cells[1].text = "30000"
    d.save(p)


def make_odt(p: Path):
    content = ('<?xml version="1.0"?><office:document-content xmlns:office="o" xmlns:text="t"><office:body><office:text>'
               '<text:h>บันทึกประชุม</text:h><text:p>ข้อสรุป &amp; มติ<text:line-break/>บรรทัดสอง</text:p>'
               '</office:text></office:body></office:document-content>')
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        z.writestr("content.xml", content)


def make_epub(p: Path):
    opf = ('<package><manifest><item id="c2" href="ch2.xhtml"/><item id="c1" href="ch1.xhtml"/></manifest>'
           '<spine><itemref idref="c1"/><itemref idref="c2"/></spine></package>')
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("OEBPS/content.opf", opf)
        z.writestr("OEBPS/ch1.xhtml", "<html><body><h1>บทที่ 1</h1><p>เริ่มต้นเรื่อง</p></body></html>")
        z.writestr("OEBPS/ch2.xhtml", "<html><body><p>ตอนจบ lighthouse</p></body></html>")


CASES = {
    "budget.xlsx": (make_xlsx, ["# งบประมาณ", "ค่าเช่าสำนักงาน\t45000"]),
    "deck.pptx": (make_pptx, ["แผนการตลาด Q3", "เปิดตัวสินค้า quokka", "speaker note ลับ"]),
    "contract.docx": (make_docx_with_table, ["สัญญาจ้าง", "ค่าจ้างรายเดือน\t30000"]),
    "minutes.odt": (make_odt, ["บันทึกประชุม", "ข้อสรุป & มติ\nบรรทัดสอง"]),
    "book.epub": (make_epub, ["บทที่ 1", "ตอนจบ lighthouse"]),
}


@pytest.mark.parametrize("name", CASES)
def test_office_and_document_formats(tmp_path, name):
    maker, expected = CASES[name]
    path = tmp_path / name
    maker(path)
    text = extract_text(path)
    for piece in expected:
        assert piece in text
    if name == "book.epub":  # follows the reading order, not file order
        assert text.index("บทที่ 1") < text.index("lighthouse")


def test_text_like_formats(tmp_path):
    files = {
        "page.html": ("<html><head><style>.x{}</style><script>evil()</script></head><body><p>สวัสดี&nbsp;โลก</p><br>บรรทัด</body></html>", ["สวัสดี\xa0โลก", "บรรทัด"], ["evil", ".x{}"]),
        "note.rtf": ("{\\rtf1\\ansi\\ansicpg874{\\fonttbl{\\f0 Tahoma;}}\\f0 Hello \\'a4\\'e8\\'d2\\'e0\\'aa\\'e8\\'d2\\par next line \\u3585?}", ["Hello ค่าเช่า\nnext line ก"], ["Tahoma"]),
        "mail.eml": ("Subject: =?utf-8?b?4LmD4Lia4LmA4Liq4LiZ4Lit4Lij4Liy4LiE4Liy?=\nFrom: a@example.com\nTo: b@example.com\nContent-Type: text/plain; charset=utf-8\n\nตกลงตามนี้ครับ\n", ["Subject: ใบเสนอราคา", "From: a@example.com", "ตกลงตามนี้ครับ"], []),
        "nb.ipynb": (json.dumps({"cells": [{"cell_type": "code", "source": ["print('สวัสดี')\n"], "outputs": [{"text": ["สวัสดี\n"]}]}]}), ["print('สวัสดี')", "สวัสดี"], []),
        "Dockerfile": ("FROM python:3.13\nRUN pip install fastapi\n", ["pip install fastapi"], []),
        "query.graphql": ("query { invoices { total } }", ["invoices"], []),
        "notes.weirdext": ("ข้อความธรรมดา unknown extension", ["unknown extension"], []),
    }
    for name, (content, wanted, unwanted) in files.items():
        p = tmp_path / name
        p.write_bytes(content.encode("utf-8")) if name != "note.rtf" else p.write_text(content, encoding="latin-1")
        text = extract_text(p)
        for w in wanted:
            assert w in text, (name, text)
        for u in unwanted:
            assert u not in text, (name, text)


def test_zip_lists_members_and_binary_is_skipped(tmp_path):
    with zipfile.ZipFile(tmp_path / "photos.zip", "w") as z:
        z.writestr("trip/ภูเก็ต-001.jpg", b"\xff\xd8")
        z.writestr("trip/", b"")
    assert extract_text(tmp_path / "photos.zip") == "trip/ภูเก็ต-001.jpg"
    (tmp_path / "blob.bin").write_bytes(b"\x00\x01binary")
    assert extract_text(tmp_path / "blob.bin") == ""
    (tmp_path / "broken.xlsx").write_bytes(b"not a zip")
    assert extract_text(tmp_path / "broken.xlsx") == ""  # corrupt files never crash a scan


def test_upgrade_reads_old_files_again(tmp_path):
    root = tmp_path / "r"
    root.mkdir()
    make_xlsx(root / "budget.xlsx")
    db = Database(tmp_path / "t.db")
    ix = Indexer(db, [str(root)])
    # Simulate an index built by the previous version (no xlsx support).
    db.set_meta("extractor_version", 1)
    st = (root / "budget.xlsx").stat()
    db.upsert_file(str(root / "budget.xlsx"), "budget.xlsx", "", st.st_mtime, st.st_size)
    assert db.search("ค่าเช่าสำนักงาน") == []
    assert ix.run().added_or_updated == 1  # unchanged on disk, but read again once
    assert [r["title"] for r in db.search("ค่าเช่าสำนักงาน")] == ["budget.xlsx"]
    assert ix.run().added_or_updated == 0  # …and only once
